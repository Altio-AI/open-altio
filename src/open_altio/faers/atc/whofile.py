"""Read WHO's ATC classification index with DDDs, the spreadsheet WHO sells.

    from open_altio.faers.atc.whofile import read_index
    codes = read_index(folder)        # [(code, name), ...] in the file's order

WHO's file is never copied into the repository or a release; only the class
table built from it is (src/open_altio/faers/atc/table.py). This module is the
one place that knows the spreadsheet's layout, so a file laid out differently
is fitted by changing `HEADERS` or `_cells` here and nothing else.

The 2026 file's first sheet, "ATC index with DDDs 2026", has one heading row,
"ATC code", "ATC level name", "DDD", "Unit", "Adm.R", "Comment", then one row
per code at every level (1, 3, 4, 5 or 7 characters) with WHO's name for it,
and for a substance its defined daily dose, unit, administration route and
comment. A substance with a dose for more than one route repeats its code and
name on a row per route. Only codes and names are read.
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook

# The headings of the columns read, in lower case
HEADERS = {"code": "atc code", "name": "atc level name"}
SUFFIX = ".xlsx"
CODE = re.compile(r"[A-Z]([0-9]{2}([A-Z]([A-Z]([0-9]{2})?)?)?)?")
LEVELS = {1: 1, 3: 2, 4: 3, 5: 4, 7: 5}  # code length -> ATC level
# WHO's cells carry stray and non-breaking spaces (which \s matches) and zero-width spaces
_SPACE = re.compile(r"\s+")
_ZERO_WIDTH = "\u200b"


class WhoRefused(Exception):
    """The folder does not hold WHO's index as this reader knows it; the message says why."""


def level(code: str) -> int:
    return LEVELS[len(code)]


def _clean(cell: object) -> str:
    return _SPACE.sub(" ", str(cell).replace(_ZERO_WIDTH, "")).strip() if cell is not None else ""


def index_file(folder: str | Path) -> Path:
    """The one spreadsheet in `folder` (Excel's own lock files aside)."""
    root = Path(folder).expanduser()
    if not root.is_dir():
        raise WhoRefused(f"{root} is not a folder")
    found = sorted(
        p for p in root.iterdir() if p.suffix.lower() == SUFFIX and not p.name.startswith("~$")
    )
    if len(found) != 1:
        names = ", ".join(p.name for p in found) or "none"
        raise WhoRefused(f"{root} must hold exactly one {SUFFIX} file, WHO's index; found {names}")
    return found[0]


def _cells(path: Path) -> Iterator[list[str]]:
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    for row in sheet.iter_rows(values_only=True):
        yield [_clean(c) for c in row]


def read_index(folder: str | Path) -> list[tuple[str, str]]:
    """Every (code, name) in WHO's index, once each, in the file's order.
    Raises WhoRefused when the file is missing, has no heading row naming the
    code and name columns, holds a code that is not an ATC code, or names one
    code two ways."""
    path = index_file(folder)
    rows = _cells(path)
    for cells in rows:
        lowered = [c.lower() for c in cells]
        if all(h in lowered for h in HEADERS.values()):
            at = {k: lowered.index(h) for k, h in HEADERS.items()}
            break
    else:
        wanted = ", ".join(repr(h) for h in HEADERS.values())
        raise WhoRefused(f"{path.name} has no heading row with the columns {wanted}")
    found: dict[str, str] = {}
    for n, cells in enumerate(rows, start=2):
        code, name = cells[at["code"]], cells[at["name"]]
        if not CODE.fullmatch(code):
            raise WhoRefused(f"{path.name}: {code!r} near row {n} is not an ATC code")
        if found.setdefault(code, name) != name:
            raise WhoRefused(f"{path.name}: {code} is named both {found[code]!r} and {name!r}")
    if not found:
        raise WhoRefused(f"{path.name} holds no ATC codes")
    return list(found.items())


__all__ = ["CODE", "HEADERS", "LEVELS", "WhoRefused", "index_file", "level", "read_index"]
